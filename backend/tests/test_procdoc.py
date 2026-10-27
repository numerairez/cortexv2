"""Process document (Word / PDF) reading and document → workflow translation."""

import io
import json
import zipfile
import zlib

import pytest

from app.connectors import procdoc, vba


def _png(w: int, h: int, colour=(200, 30, 30)) -> bytes:
    row = bytes(colour) * w
    return procdoc._png(w, h, row * h, w * 3)


def _shape_xml(texts: list[str]) -> str:
    """A floating Word drawing made of one text box, like a box in a hand-drawn flowchart."""
    paras = "".join(f"<w:p><w:r><w:t>{t}</w:t></w:r></w:p>" for t in texts)
    return (
        '<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
        'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
        'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape" '
        'xmlns:v="urn:schemas-microsoft-com:vml">'
        '<mc:AlternateContent><mc:Choice Requires="wps"><w:drawing><wp:anchor>'
        '<wp:docPr id="9" name="Flowchart 1"/>'
        '<a:graphic><a:graphicData uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">'
        f'<wps:wsp><wps:txbx><w:txbxContent>{paras}</w:txbxContent></wps:txbx></wps:wsp>'
        '</a:graphicData></a:graphic></wp:anchor></w:drawing></mc:Choice>'
        f'<mc:Fallback><w:pict><v:shape><v:textbox><w:txbxContent>{paras}</w:txbxContent></v:textbox></v:shape></w:pict></mc:Fallback>'
        '</mc:AlternateContent></w:r>'
    )


@pytest.fixture()
def procedure_docx(tmp_path):
    import docx
    from docx.oxml import parse_xml

    d = docx.Document()
    d.add_heading("Month-end sales summary", 0)
    d.add_paragraph("Open S:\\Finance\\sales_jan.xlsx and the product list \"product master.xlsx\".")
    d.add_heading("Steps", 1)
    d.add_paragraph("Remove rows where Qty is zero.", style="List Bullet")
    d.add_paragraph("Email the summary to the branch manager for approval.", style="List Bullet")
    big = tmp_path / "screen.png"
    big.write_bytes(_png(300, 200))
    d.add_picture(str(big))
    d.inline_shapes[0]._inline.docPr.set("descr", "Pivot settings: rows Region, values Sum of Qty")
    d.add_picture(str(big))                                  # the same picture again
    icon = tmp_path / "icon.png"
    icon.write_bytes(_png(16, 16, (0, 0, 0)))
    d.add_picture(str(icon))
    p = d.add_paragraph()
    p._p.append(parse_xml(_shape_xml(["Start", "Filter Region = Luzon"])))
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Column", "Rule"
    t.cell(1, 0).text, t.cell(1, 1).text = "Qty", "must be > 0"
    d.add_paragraph("Done.")
    path = tmp_path / "Procedure.docx"
    d.save(str(path))
    return path


def _pdf(pages: list[bytes], image: tuple[int, int, bytes] | None = None) -> bytes:
    """A minimal PDF: one content stream per page, Helvetica, and optionally one RGB image as /Im1."""
    objs: list[bytes] = []
    kids = []
    font = 3 + 2 * len(pages)
    img = font + 1
    for i, content in enumerate(pages):
        page_no, stream_no = 3 + 2 * i, 4 + 2 * i
        kids.append(f"{page_no} 0 R")
        res = f"/Font << /F1 {font} 0 R >>" + (f" /XObject << /Im1 {img} 0 R >>" if image else "")
        objs.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << {res} >> "
                    f"/Contents {stream_no} 0 R >>".encode())
        objs.append(f"<< /Length {len(content)} >>\nstream\n".encode() + content + b"\nendstream")
    objs.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    if image:
        w, h, rgb = image
        data = zlib.compress(rgb)
        objs.append(f"<< /Type /XObject /Subtype /Image /Width {w} /Height {h} /ColorSpace /DeviceRGB "
                    f"/BitsPerComponent 8 /Filter /FlateDecode /Length {len(data)} >>\nstream\n".encode() + data + b"\nendstream")
    objs = [b"<< /Type /Catalog /Pages 2 0 R >>",
            f"<< /Type /Pages /Kids [{' '.join(kids)}] /Count {len(pages)} >>".encode()] + objs
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for n, body in enumerate(objs, start=1):
        offsets.append(len(out))
        out += f"{n} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    out += b"".join(f"{o:010d} 00000 n \n".encode() for o in offsets)
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    return bytes(out)


def _text(s: str, y: int = 720) -> bytes:
    return f"BT /F1 12 Tf 72 {y} Td ({s}) Tj ET\n".encode()


@pytest.fixture()
def procedure_pdf(tmp_path):
    lines = b"".join(f"{72 + i * 10} 400 m {72 + i * 10} 300 l S\n".encode() for i in range(25))   # a vector drawing
    pages = [
        _text("Step 1: open targets.csv and keep Luzon rows."),
        _text("Step 2: follow the flowchart below.") + lines,
        b"q 500 0 0 400 50 200 cm /Im1 Do Q\n",                                  # a scan: a big picture, no text
    ]
    path = tmp_path / "Procedure.pdf"
    path.write_bytes(_pdf(pages, image=(40, 30, bytes([90, 90, 200]) * 40 * 30)))
    return path


def test_docx_text_and_figures(data_dir, procedure_docx):
    r = procdoc.load(procedure_docx)
    code = r["modules"][0]["code"]
    lines = code.split("\n")
    assert lines[0] == "## Month-end sales summary"
    assert "- Remove rows where Qty is zero." in lines
    assert "| Qty | must be > 0 |" in lines
    assert lines.index("| Qty | must be > 0 |") < lines.index("Done.")             # tables stay in reading order
    assert [h["name"] for h in r["modules"][0]["procedures"]] == ["Month-end sales summary", "Steps"]

    figs = r["figures"]
    assert [f["kind"] for f in figs] == ["image", "image", "drawing"]               # the repeat is not a new figure
    pivot, icon, drawing = figs
    assert pivot["alt"] == "Pivot settings: rows Region, values Sum of Qty"
    assert pivot["mode"] == "attach" and pivot["file"] == "fig-01.png" and (pivot["width"], pivot["height"]) == (300, 200)
    assert len(pivot["lines"]) == 2 and "[Figure 1 again]" in lines
    assert lines[pivot["line"] - 1].startswith("[Figure 1: picture · alt text")
    assert icon["mode"] == "skip"
    assert drawing["mode"] == "describe" and drawing["file"] is None
    assert drawing["text"] == "Start / Filter Region = Luzon"                      # read once, not again from the fallback
    assert not any("Start" in ln for ln in lines if not ln.startswith("["))      # shape text isn't body text

    deps = {d["filename"]: d for d in r["dependencies"]}
    assert deps["sales_jan.xlsx"]["reference"] == "S:\\Finance\\sales_jan.xlsx"
    assert "product master.xlsx" in deps
    assert {a["action"] for a in r["actions"]} >= {"Sends email", "Needs approval or sign-off"}
    saved = procdoc.figure_path(r["figures_key"], "fig-01.png").read_bytes()
    assert procdoc.image_info(saved) == ("png", 300, 200)
    with pytest.raises(vba.VBAError):
        procdoc.figure_path(r["figures_key"], "../../workbench.sqlite3")


def test_pdf_pages_become_pictures_when_needed(data_dir, procedure_pdf):
    r = procdoc.load(procedure_pdf)
    code = r["modules"][0]["code"]
    assert "Step 1: open targets.csv" in code
    pages = {f["locator"]: f for f in r["figures"]}
    assert set(pages) == {"p. 2", "p. 3"}                                           # page 1 is text only
    assert pages["p. 2"]["reason"] == "has a drawing or bordered table"
    assert pages["p. 3"]["reason"] == "scanned page"
    fmt, w, h = procdoc.image_info(procdoc.figure_path(r["figures_key"], pages["p. 2"]["file"]).read_bytes())
    assert fmt == "png" and h > w > 1000
    assert r["dependencies"][0]["filename"] == "targets.csv"


def test_text_only_procedure_keeps_step_numbers(data_dir, tmp_path):
    """Most procedures are words only: numbered steps (as Word numbers them) and no pictures."""
    import docx

    d = docx.Document()
    d.add_heading("Daily cash reconciliation", 1)
    for step in ["Download the bank statement as bank.csv.", "Remove rows where Amount is blank.",
                 "Match each row to the ledger by Reference.", "Repeat step 3 for each branch."]:
        d.add_paragraph(step, style="List Number")
    d.add_paragraph("Notes:")
    d.add_paragraph("Ignore reversals.", style="List Bullet")
    path = tmp_path / "Cash.docx"
    d.save(str(path))

    r = procdoc.load(path)
    lines = r["modules"][0]["code"].split("\n")
    assert lines[1:5] == ["1. Download the bank statement as bank.csv.", "2. Remove rows where Amount is blank.",
                          "3. Match each row to the ledger by Reference.", "4. Repeat step 3 for each branch."]
    assert lines[-1] == "- Ignore reversals."
    assert r["figures"] == [] and r["figures_key"] is None
    assert r["dependencies"][0]["filename"] == "bank.csv"


def test_word_numbering_levels_and_formats():
    from docx.oxml import parse_xml

    ns = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
    root = parse_xml(f"""<w:numbering {ns}>
      <w:abstractNum w:abstractNumId="0">
        <w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="decimal"/><w:lvlText w:val="Step %1:"/></w:lvl>
        <w:lvl w:ilvl="1"><w:start w:val="1"/><w:numFmt w:val="lowerLetter"/><w:lvlText w:val="%1.%2)"/></w:lvl>
      </w:abstractNum>
      <w:num w:numId="5"><w:abstractNumId w:val="0"/></w:num>
      <w:num w:numId="6"><w:abstractNumId w:val="0"/><w:lvlOverride w:ilvl="0"><w:startOverride w:val="7"/></w:lvlOverride></w:num>
    </w:numbering>""")

    class Doc:
        styles = []
        part = type("P", (), {"numbering_part": type("N", (), {"element": root})()})()

    n = procdoc._Numbering(Doc())

    def ppr(num, lvl):
        return parse_xml(f'<w:pPr {ns}><w:numPr><w:ilvl w:val="{lvl}"/><w:numId w:val="{num}"/></w:numPr></w:pPr>')

    got = [n.label(ppr(5, 0), ""), n.label(ppr(5, 1), ""), n.label(ppr(5, 1), ""), n.label(ppr(5, 0), ""),
           n.label(ppr(5, 1), ""), n.label(ppr(6, 0), "")]
    assert got == ["Step 1:", "1.a)", "1.b)", "Step 2:", "2.a)", "Step 7:"]
    assert n.label(None, "") is None


def test_unreadable_documents(tmp_path):
    old = tmp_path / "old.doc"
    old.write_bytes(b"\xd0\xcf\x11\xe0")
    with pytest.raises(vba.VBAError, match="save it as .docx"):
        procdoc.load(old)
    with pytest.raises(vba.VBAError, match="paste its text"):
        procdoc.load()


def test_doc_to_workflow_end_to_end(client, files, procedure_docx):
    a = client.post("/api/macros/analyse", json={"path": str(procedure_docx), "language": "doc"}).json()
    figs = a["figures"]
    figs[2]["description"] = "Arrow from Start to the filter box"
    assert client.get(f"/api/figures/{a['figures_key']}/fig-01.png").content[:4] == b"\x89PNG"
    assert client.get(f"/api/figures/{a['figures_key']}/secret.txt").status_code == 404

    src = client.post("/api/sources", json={"path": str(files["jan"])}).json()
    sel = client.post("/api/selections", json={"source_id": src["id"], "label": "Sales", "sheet": "Sales", "header_row": 5}).json()
    run = client.post("/api/runs", json={
        "service_id": "doc_to_workflow", "bindings": {"sales": sel["id"]},
        "options": {"macro_name": "Procedure", "entry": "Steps", "modules": a["modules"], "dependencies": [],
                    "actions": a["actions"], "figures": figs, "figures_key": a["figures_key"]}}).json()
    prompt = run["state"]["prompt"]
    assert "written business procedure" in prompt and "# Procedure document" in prompt and "VBA" not in prompt
    assert "- Figure 1 (picture, line" in prompt and "attached as fig-01.png" in prompt
    assert "- Figure 2 (picture" in prompt and prompt.split("- Figure 2", 1)[1].split("\n", 1)[0].endswith("not provided")
    assert "the user describes it as: Arrow from Start" in prompt
    assert '"doc_lines": "Procedure.docx:12-20"' in prompt
    assert run["state"]["figures_attached"] == 1

    z = zipfile.ZipFile(io.BytesIO(client.get(f"/api/runs/{run['id']}/figures.zip").content))
    assert z.namelist() == ["fig-01.png"]

    reply = {"recipe": {"inputs": [{"role": "sales"}], "base": "sales",
                        "steps": [{"op": "filter_rows", "conditions": [{"column": "Qty", "operator": "gt", "value": 0}]}]},
             "step_notes": [{"step": 1, "doc_lines": "Procedure.docx:5", "explanation": "Drops zero quantities"}],
             "not_translated": [{"doc_lines": "Procedure.docx:6", "what_it_does": "Emails the manager", "why": "manual"}]}
    done = client.post(f"/api/runs/{run['id']}/response", json={"text": json.dumps(reply)}).json()
    assert done["status"] == "completed", done["errors"]
    assert done["state"]["recipe"]["steps"][0]["note"] == "From document Procedure.docx:5: Drops zero quantities"
    assert done["state"]["not_translated"][0]["macro_lines"] == "Procedure.docx:6"
    assert done["state"]["recipe"]["groups"][0]["source"]["language"] == "doc"
