from app.connectors import tabular
from app.transformations import export


def test_preview_uses_physical_excel_rows(files):
    p = tabular.preview_rows(files["jan"], "Sales", 10)
    assert [r["row"] for r in p["rows"]][:6] == [1, 2, 3, 4, 5, 6]
    assert p["rows"][4]["cells"][0] == "Txn ID"
    assert p["suggested_header_row"] == 5


def test_load_with_header_on_row_5(files):
    t = tabular.load_table(files["jan"], "Sales", header_row=5)
    assert list(t.df.columns) == ["Txn ID", "Date", "Region", "Product ID", "Qty", "Unit Price", "Note"]
    assert len(t.df) == 5
    assert t.first_data_row == 6
    types = {c["name"]: c["type"] for c in tabular.schema_of(t.df)}
    assert types["Txn ID"] == "integer" and types["Date"] == "date" and types["Unit Price"] == "number"


def test_wrong_header_row_produces_different_schema(files):
    t = tabular.load_table(files["jan"], "Sales", header_row=1)
    assert "Txn ID" not in t.df.columns
    drift = tabular.schema_drift(tabular.schema_of(tabular.load_table(files["jan"], "Sales", 5).df),
                                 tabular.schema_of(t.df))
    assert drift and "Txn ID" in drift["missing"]


def test_csv_single_sheet(files):
    assert tabular.list_sheets(files["csv"]) == [tabular.CSV_SHEET]
    t = tabular.load_table(files["csv"], tabular.CSV_SHEET, 1)
    assert list(t.df["Region"]) == ["Luzon", "Visayas"]


def test_xlsx_export_never_writes_formulas(files, tmp_path):
    from openpyxl import load_workbook

    t = tabular.load_table(files["jan"], "Sales", 5)
    out = tmp_path / "out.xlsx"
    export.write_xlsx(t.df, out, [("Note", "=1+1")])
    ws = load_workbook(out)["Result"]
    cells = [c for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("=")]
    assert cells and all(c.data_type == "s" for c in cells)


def test_csv_export_escapes_formula_text(files, tmp_path):
    t = tabular.load_table(files["jan"], "Sales", 5)
    out = tmp_path / "out.csv"
    export.write_csv(t.df, out)
    text = out.read_text(encoding="utf-8-sig")
    assert "'=HYPERLINK" in text
    assert ",=HYPERLINK" not in text


def test_column_values_match_per_value_conversion():
    import datetime as dt
    import math

    import pandas as pd

    df = tabular.normalise_types(pd.DataFrame({
        "when": [dt.datetime(2024, 1, 5), None, dt.datetime(2024, 2, 1, 9, 30)],
        "precise": [dt.datetime(2024, 1, 5, 0, 0, 0, 1500), dt.datetime(2024, 1, 6), None],
        "n": [1, None, 3],
        "x": [1.5, float("inf"), None],
        "flag": [True, None, False],
        "name": ["a", None, "c"],
    }, dtype=object))
    for col in df.columns:
        expected = [tabular.to_json_value(v) for v in df[col]]
        expected = [None if isinstance(v, float) and math.isinf(v) else v for v in expected]
        assert tabular.column_values(df[col]) == expected, col
    assert tabular.column_values(df["n"], limit=1) == [1]


def _openpyxl_only(monkeypatch):
    monkeypatch.setattr(tabular, "_calamine", lambda path: None)


def _awkward_workbook(path):
    """Leading blank rows and columns, whole and fractional numbers, dates with and without a time,
    times, booleans, padded and blank text, a hidden sheet."""
    import datetime as dt

    from .conftest import _write_xlsx

    return _write_xlsx(path, [
        [None], [None, None, "Title"], [None],
        [None, "ID", "Amount", "When", "At", "Flag", "Note", "Big"],
        [None, 2, 1.5, dt.datetime(2026, 1, 3), dt.time(9, 15), True, " padded ", 10**15],
        [None, "P3", 3.0, dt.datetime(2026, 1, 3, 10, 30), None, False, "", -7],
        [None, None, None, None, None, None, None, None],
        [None, 4, 0.1, None, dt.time(0, 0), None, "=not a formula", 0],
    ], sheet="Data", extra_sheets={"Hidden": [["x"]]})


def test_calamine_reads_exactly_what_openpyxl_reads(files, tmp_path, monkeypatch):
    books = [(files["jan"], "Sales", 5), (files["jan"], "Notes", 1), (files["feb"], "Sales", 5),
             (files["products"], "Products", 1), (_awkward_workbook(tmp_path / "awkward.xlsx"), "Data", 4)]
    for path, sheet, header in books:
        assert tabular._calamine(path) is not None
        fast_sheets = tabular.list_sheets(path)
        fast_rows = list(tabular.iter_physical_rows(path, sheet))
        fast_preview = list(tabular.iter_physical_rows(path, sheet, start_row=2, max_rows=3))
        fast_table = tabular.load_table(path, sheet, header)
        with monkeypatch.context() as m:
            _openpyxl_only(m)
            assert tabular.list_sheets(path) == fast_sheets
            slow_rows = list(tabular.iter_physical_rows(path, sheet))
            # openpyxl pads every row to the sheet's width; compare without trailing blanks.
            assert [(n, tabular._trim(r)) for n, r in fast_rows] == [(n, tabular._trim(r)) for n, r in slow_rows], path.name
            assert [(n, tabular._trim(r)) for n, r in fast_preview] == \
                [(n, tabular._trim(r)) for n, r in tabular.iter_physical_rows(path, sheet, start_row=2, max_rows=3)]
            slow_table = tabular.load_table(path, sheet, header)
        assert fast_table.df.equals(slow_table.df) and list(fast_table.df.dtypes) == list(slow_table.df.dtypes), path.name
        assert fast_table.warnings == slow_table.warnings and fast_table.letters == slow_table.letters


def test_sheet_with_error_cells_is_read_with_openpyxl(tmp_path, monkeypatch):
    import re
    import zipfile

    plain = _awkward_workbook(tmp_path / "plain.xlsx")
    broken = tmp_path / "errors.xlsx"
    with zipfile.ZipFile(plain) as src, zipfile.ZipFile(broken, "w") as out:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":   # a cached #N/A where C5 (1.5) was
                data = re.sub(rb'<c r="C5"[^>]*>.*?</c>', b'<c r="C5" t="e"><f>NA()</f><v>#N/A</v></c>', data)
            out.writestr(item, data)

    assert not tabular._sheet_has_errors(plain, "Data") and not tabular._sheet_has_errors(broken, "Hidden")
    assert tabular._sheet_has_errors(broken, "Data")
    rows = dict(tabular.iter_physical_rows(broken, "Data"))
    assert rows[5][2] == "#N/A"   # kept visible, not read as a blank


def _csv(tmp_path, text, name="data.csv"):
    p = tmp_path / name
    p.write_text(text, encoding="utf-8")
    return p


def test_csv_columns_are_typed_only_when_unambiguous(tmp_path):
    p = _csv(tmp_path, "\n".join([
        "Qty,Price,Code,Account,When,Stamp,Local date,Flag,Mixed,Spaced",
        "2,1.50,00123,1234567890123456,2026-01-03,2026-01-03 09:15:00,03/01/2026,TRUE,7,  12 ",
        "3,2,00456,1234567890123457,2026-02-01,2026-02-01T10:30,04/01/2026,false,n/a,3",
        ",-0.25,,,,,,,,",
    ]) + "\n")
    df = tabular.load_table(p, tabular.CSV_SHEET, 1).df
    types = {c["name"]: c["type"] for c in tabular.schema_of(df)}
    assert types == {"Qty": "integer", "Price": "number", "Code": "text", "Account": "text", "When": "date",
                     "Stamp": "date", "Local date": "text", "Flag": "boolean", "Mixed": "text", "Spaced": "integer"}
    assert tabular.column_values(df["Qty"]) == [2, 3, None]
    assert tabular.column_values(df["Price"]) == [1.5, 2.0, -0.25]
    assert list(df["Code"]) == ["00123", "00456", None]
    assert tabular.column_values(df["Stamp"]) == ["2026-01-03T09:15:00", "2026-02-01T10:30:00", None]
    assert tabular.column_values(df["Flag"]) == [True, False, None]


def test_csv_fast_reader_matches_row_by_row_reader(tmp_path, monkeypatch):
    """The pandas reader gives what the csv module gives: same physical rows, cleaning and warnings."""
    files = [
        (_csv(tmp_path, "Report,,\nGenerated today,,\n,,\nRegion,Target,Note\n Luzon , 100 ,\"a, b\"\n,,\n"
                        "Visayas,50,\"line one\nline two\"\n,,\nMindanao,  ,\"  \"\n,,\n,,\n", "titled.csv"), 4),
        (_csv(tmp_path, "a;b;;d\n1;x;;\n2;y;;4\n", "semicolon.csv"), 1),
        (_csv(tmp_path, "Name\tAmount\nx\t1\ny\t\n", "tabs.tsv"), 1),
        (_csv(tmp_path, "Name,Name,\n1,2,3\n", "dupes.csv"), 1),
        (_csv(tmp_path, "Name,Amount\n", "header_only.csv"), 1),
    ]
    for path, header in files:
        fast = tabular.load_table(path, tabular.CSV_SHEET, header)
        fast_head = tabular.load_table(path, tabular.CSV_SHEET, header, max_rows=1)
        with monkeypatch.context() as m:
            m.setattr(tabular, "_csv_grid", lambda *a: None)
            slow = tabular.load_table(path, tabular.CSV_SHEET, header)
            slow_head = tabular.load_table(path, tabular.CSV_SHEET, header, max_rows=1)
        for a, b in ((fast, slow), (fast_head, slow_head)):
            assert a.df.equals(b.df) and list(a.df.dtypes) == list(b.df.dtypes), path.name
            assert a.warnings == b.warnings and a.letters == b.letters and a.original_headers == b.original_headers

    t = tabular.load_table(files[0][0], tabular.CSV_SHEET, 4)
    assert list(t.df["Region"]) == ["Luzon", "Visayas", "Mindanao"]
    assert tabular.column_values(t.df["Target"]) == [100, 50, None]
    assert list(t.df["Note"]) == ["a, b", "line one\nline two", None]
    assert "2 blank row(s) inside the table were removed" in t.warnings


def test_csv_with_rows_wider_than_the_first_is_read_row_by_row(tmp_path):
    p = _csv(tmp_path, "Sales export\nRegion,Target\nLuzon,100\n")
    assert tabular._csv_grid(p, 2, None) is None
    df = tabular.load_table(p, tabular.CSV_SHEET, 2).df
    assert list(df["Region"]) == ["Luzon"] and tabular.column_type(df["Target"]) == "integer"


def test_csv_padding_check_sees_whitespace_beside_delimiters_quotes_and_line_ends():
    assert not tabular._has_padded_cells(b"a,b c\n1,two words\r\n", ",")
    for raw in (b" a,b\n", b"a ,b\n", b"a, b\n", b"a,b \n", b"a,b \r\n", b'a," b"\n', b"a,b\n \n", b"a,b "):
        assert tabular._has_padded_cells(raw, ","), raw
    assert not tabular._has_padded_cells(b"a\t\tb\n", "\t")


def test_csv_schema_saved_as_text_does_not_drift_once_typed():
    saved = [{"name": "Qty", "type": "text"}, {"name": "When", "type": "text"}, {"name": "Code", "type": "integer"}]
    now = [{"name": "Qty", "type": "integer"}, {"name": "When", "type": "date"}, {"name": "Code", "type": "text"}]
    assert tabular.schema_drift(saved, now, from_csv=True)["type_changes"] == [{"name": "Code", "was": "integer", "now": "text"}]
    assert len(tabular.schema_drift(saved, now)["type_changes"]) == 3


def test_file_hash_is_remembered_until_the_file_changes(tmp_path, monkeypatch):
    import os

    from app.connectors import documents

    p = _csv(tmp_path, "a\n1\n")
    first = documents.fingerprint_file(p)["sha256"]
    opened = []
    real_open = type(p).open
    monkeypatch.setattr(type(p), "open", lambda self, *a, **k: opened.append(self) or real_open(self, *a, **k))
    assert documents.fingerprint_file(p)["sha256"] == first and not opened
    p.write_text("a\n2\n", encoding="utf-8")   # same size
    st = p.stat()
    os.utime(p, ns=(st.st_atime_ns, st.st_mtime_ns + 1_000_000))
    assert documents.fingerprint_file(p)["sha256"] != first and opened
