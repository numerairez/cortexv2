"""Text extraction from local documents into locatable segments.

Each segment carries a human-readable ``locator`` (page number, heading, paragraph)
so prompts can cite exactly where a passage came from.
"""

from __future__ import annotations

import hashlib
import re
from dataclasses import dataclass
from html.parser import HTMLParser
from pathlib import Path

TEXT_EXTS = {".txt", ".md", ".markdown"}
DOC_EXTS = TEXT_EXTS | {".pdf", ".docx"}
MAX_SEGMENT_CHARS = 4000


class DocumentError(Exception):
    pass


@dataclass
class Segment:
    locator: str
    text: str


def fingerprint_file(path: Path) -> dict:
    h = hashlib.sha256()
    try:
        with path.open("rb") as f:
            for chunk in iter(lambda: f.read(1 << 20), b""):
                h.update(chunk)
        st = path.stat()
    except FileNotFoundError:
        raise DocumentError("File not found")
    except PermissionError:
        raise DocumentError("Permission denied (the file may be open in another program)")
    return {"sha256": h.hexdigest(), "size": st.st_size, "mtime": st.st_mtime}


def extract_segments(path: Path) -> list[Segment]:
    ext = path.suffix.lower()
    try:
        if ext in {".md", ".markdown"}:
            return _split_long(_markdown_segments(_read_text(path)))
        if ext == ".txt":
            return _split_long(_paragraph_segments(_read_text(path)))
        if ext == ".pdf":
            return _split_long(_pdf_segments(path))
        if ext == ".docx":
            return _split_long(_docx_segments(path))
    except DocumentError:
        raise
    except Exception as e:
        raise DocumentError(f"Could not read {path.name}: {e.__class__.__name__}: {e}") from e
    raise DocumentError(f"Unsupported document type: {ext}")


def _read_text(path: Path) -> str:
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "cp1252"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            pass
    return raw.decode("latin-1")


def _paragraph_segments(text: str) -> list[Segment]:
    paras = [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]
    return [Segment(f"para {i}", p) for i, p in enumerate(paras, start=1)]


def _markdown_segments(text: str) -> list[Segment]:
    segs: list[Segment] = []
    heading = "Introduction"
    buf: list[str] = []

    def flush():
        body = "\n".join(buf).strip()
        if body:
            segs.append(Segment(f"§ {heading}", body))

    for line in text.splitlines():
        m = re.match(r"^(#{1,6})\s+(.*)$", line)
        if m:
            flush()
            buf = []
            heading = m.group(2).strip()
        else:
            buf.append(line)
    flush()
    return segs


def _pdf_segments(path: Path) -> list[Segment]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    if reader.is_encrypted:
        raise DocumentError("PDF is password-protected")
    segs = []
    for i, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if text:
            segs.append(Segment(f"p. {i}", text))
    if not segs:
        raise DocumentError("No extractable text (the PDF may be scanned images)")
    return segs


def _docx_segments(path: Path) -> list[Segment]:
    import docx

    d = docx.Document(str(path))
    segs: list[Segment] = []
    heading = None
    buf: list[str] = []
    para_no = 0

    def flush():
        body = "\n".join(buf).strip()
        if body:
            segs.append(Segment(f"§ {heading}" if heading else f"para {para_no}", body))

    for p in d.paragraphs:
        para_no += 1
        style = (p.style.name or "").lower() if p.style is not None else ""
        if style.startswith("heading") or style == "title":
            flush()
            buf = []
            heading = p.text.strip() or heading
        elif p.text.strip():
            buf.append(p.text.strip())
    flush()
    for ti, table in enumerate(d.tables, start=1):
        rows = [" | ".join(c.text.strip() for c in r.cells) for r in table.rows]
        if rows:
            segs.append(Segment(f"table {ti}", "\n".join(rows)))
    return segs


def _split_long(segs: list[Segment]) -> list[Segment]:
    out = []
    for s in segs:
        if len(s.text) <= MAX_SEGMENT_CHARS:
            out.append(s)
            continue
        parts = [s.text[i : i + MAX_SEGMENT_CHARS] for i in range(0, len(s.text), MAX_SEGMENT_CHARS)]
        out.extend(Segment(f"{s.locator} (part {n})", t) for n, t in enumerate(parts, start=1))
    return out


class _HTMLText(HTMLParser):
    """Turn Confluence storage-format HTML into heading-anchored segments."""

    BLOCK = {"p", "li", "tr", "br", "div", "pre", "blockquote"}

    def __init__(self):
        super().__init__()
        self.segments: list[Segment] = []
        self.heading = "Introduction"
        self.buf: list[str] = []
        self.in_heading = False
        self.heading_buf: list[str] = []
        self.skip = 0

    def handle_starttag(self, tag, attrs):
        if tag in {"script", "style"}:
            self.skip += 1
        elif re.fullmatch(r"h[1-6]", tag):
            self._flush()
            self.in_heading = True
            self.heading_buf = []
        elif tag in self.BLOCK:
            self.buf.append("\n")
        elif tag in {"td", "th"}:
            self.buf.append(" | ")

    def handle_endtag(self, tag):
        if tag in {"script", "style"}:
            self.skip = max(0, self.skip - 1)
        elif re.fullmatch(r"h[1-6]", tag):
            self.in_heading = False
            self.heading = " ".join("".join(self.heading_buf).split()) or self.heading

    def handle_data(self, data):
        if self.skip:
            return
        (self.heading_buf if self.in_heading else self.buf).append(data)

    def _flush(self):
        text = re.sub(r"[ \t]+", " ", "".join(self.buf))
        text = re.sub(r"\n\s*\n+", "\n", text).strip()
        if text:
            self.segments.append(Segment(f"§ {self.heading}", text))
        self.buf = []

    def close(self):
        super().close()
        self._flush()


def html_segments(html: str) -> list[Segment]:
    p = _HTMLText()
    p.feed(html)
    p.close()
    return _split_long(p.segments)
