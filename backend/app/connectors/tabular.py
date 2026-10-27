"""Excel/CSV reading: sheet listing, physical-row previews, header detection, profiling.

Row numbers exposed to users are always 1-based physical worksheet rows, exactly as
Excel shows them. We read rows through calamine/openpyxl/csv directly instead of
``pandas.read_excel(header=...)`` because pandas may drop leading blank rows, which
would silently shift a user's "header is on row 5" choice.

Workbooks are read with python-calamine when it is installed (about 8x faster than
openpyxl on large sheets), its values converted to what openpyxl gives. Calamine reads
error cells (#N/A, #DIV/0!) as blanks, so a sheet that contains any is read with
openpyxl instead, which keeps the error text visible.
"""

from __future__ import annotations

import csv
import datetime as dt
import io
import re
import xml.etree.ElementTree as ET
import zipfile
from dataclasses import dataclass, field, replace
from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd

EXCEL_EXTS = {".xlsx", ".xlsm"}
CSV_EXTS = {".csv", ".tsv", ".txt"}
CSV_SHEET = "(csv)"
MAX_PREVIEW_ROWS = 200


class TabularError(Exception):
    pass


def is_tabular(path: Path) -> bool:
    return path.suffix.lower() in EXCEL_EXTS | {".csv", ".tsv"}


def _open_workbook(path: Path):
    import openpyxl

    try:
        # data_only: read cached formula results, never formulas themselves.
        return openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as e:  # password-protected / corrupt / locked files land here
        raise TabularError(f"Could not open workbook: {e.__class__.__name__}: {e}") from e


def _calamine(path: Path):
    """The workbook opened with python-calamine, or None to read it with openpyxl instead
    (not installed, or a file calamine can't read that openpyxl might)."""
    try:
        from python_calamine import CalamineWorkbook
    except ImportError:
        return None
    try:
        return CalamineWorkbook.from_path(str(path))
    except Exception:
        return None


_NS_MAIN = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
_NS_REL = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
_NS_PKG = "{http://schemas.openxmlformats.org/package/2006/relationships}"


def _sheet_has_errors(path: Path, sheet: str) -> bool:
    """Whether the sheet holds any error cell (t="e"). Unsure counts as yes."""
    try:
        with zipfile.ZipFile(path) as z:
            rid = next((s.get(f"{_NS_REL}id") for s in ET.fromstring(z.read("xl/workbook.xml")).iter(f"{_NS_MAIN}sheet")
                        if s.get("name") == sheet), None)
            target = next((r.get("Target") for r in ET.fromstring(z.read("xl/_rels/workbook.xml.rels")).iter(f"{_NS_PKG}Relationship")
                           if r.get("Id") == rid), None)
            if not target:
                return True
            part = target.lstrip("/") if target.startswith("/") else f"xl/{target}"
            with z.open(part) as f:
                tail = b""
                while chunk := f.read(1 << 20):
                    window = tail + chunk
                    if b't="e"' in window or b"t='e'" in window:
                        return True
                    tail = chunk[-6:]
    except (KeyError, OSError, zipfile.BadZipFile, ET.ParseError):
        return True
    return False


def _clean_calamine_cell(v: Any) -> Any:
    """``_clean_cell`` of the value openpyxl would report: whole numbers as int (Excel stores
    3 as "3", which openpyxl reads as an int) and dates as datetimes. Runs once per cell, so
    it dispatches on the exact type."""
    t = type(v)
    if t is str:
        v = v.strip()
        return v if v != "" else None
    if t is float:
        return int(v) if v.is_integer() else v
    if t is dt.date:
        return dt.datetime(v.year, v.month, v.day)
    return v


def _calamine_rows(path: Path, sheet: str, start_row: int, max_rows: int | None) -> list[list] | None:
    """Physical rows from ``start_row`` read with calamine, or None to read with openpyxl."""
    wb = _calamine(path)
    if wb is None or sheet not in wb.sheet_names or _sheet_has_errors(path, sheet):
        return None
    try:
        rows = wb.get_sheet_by_name(sheet).to_python(
            skip_empty_area=False, nrows=start_row + max_rows - 1 if max_rows else None)
    except Exception:
        return None
    return rows[start_row - 1:]


def list_sheets(path: Path) -> list[str]:
    ext = path.suffix.lower()
    if ext in EXCEL_EXTS:
        cwb = _calamine(path)
        if cwb is not None:
            from python_calamine import SheetTypeEnum

            # Worksheets only, hidden ones included, as openpyxl's wb.worksheets lists them.
            return [m.name for m in cwb.sheets_metadata if m.typ == SheetTypeEnum.WorkSheet]
        wb = _open_workbook(path)
        try:
            return [ws.title for ws in wb.worksheets]
        finally:
            wb.close()
    if ext in CSV_EXTS:
        return [CSV_SHEET]
    if ext == ".xls":
        raise TabularError("Legacy .xls workbooks are not supported; please save as .xlsx.")
    raise TabularError(f"Unsupported tabular file type: {ext}")


def _read_csv_text(path: Path) -> str:
    return _decode_csv(path.read_bytes())[0]


def _decode_csv(raw: bytes) -> tuple[str, str]:
    """The text and the encoding it was decoded with."""
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return raw.decode(enc), enc
        except UnicodeDecodeError:
            continue
    raise TabularError("Could not decode CSV text")


def _csv_delimiter(path: Path, text: str) -> str:
    if path.suffix.lower() == ".tsv":
        return "\t"
    try:
        return csv.Sniffer().sniff(text[:4096], delimiters=",;\t|").delimiter
    except csv.Error:
        return ","


def _csv_grid(path: Path, start_row: int, max_rows: int | None) -> pd.DataFrame | None:
    """CSV rows from 1-based ``start_row`` as an object frame, cells cleaned as ``_clean_cell`` does.
    Read with pandas' C parser, several times faster than the csv module on large files; rows are
    counted as the csv module counts them, so physical row numbers agree. None when a row is wider
    than the first one read, which the C parser rejects: the caller then reads row by row."""
    raw = path.read_bytes()
    text, encoding = _decode_csv(raw)
    delim = _csv_delimiter(path, text)
    del text
    try:
        grid = pd.read_csv(io.BytesIO(raw), encoding=encoding, sep=delim, header=None, dtype=object,
                           keep_default_na=False, na_values=[""], skip_blank_lines=False, engine="c",
                           nrows=start_row - 1 + max_rows if max_rows else None)
    except pd.errors.EmptyDataError:
        return pd.DataFrame(dtype=object)
    except pd.errors.ParserError:
        return None
    grid = grid.iloc[start_row - 1:].reset_index(drop=True)
    if _has_padded_cells(raw, delim):
        for c in grid.columns:
            if grid[c].notna().any():
                s = grid[c].str.strip()
                grid[c] = s.where(s != "", None)
    return grid


def _has_padded_cells(raw: bytes, delim: str) -> bool:
    """Whether a space or tab sits next to a delimiter, quote or line break anywhere in the file;
    if not, no cell has whitespace to strip, and stripping (a Python call per cell) is skipped.
    Scanned as bytes, which every encoding we decode agrees on for these characters."""
    space = [32] if delim == "\t" else [32, 9]
    edge = [ord(delim), 34, 10, 13]
    chunk = 1 << 24
    for start in range(0, len(raw), chunk):
        # One byte of overlap, so a pair split across chunks is still seen.
        a = np.frombuffer(raw, np.uint8, count=min(chunk + 1, len(raw) - start), offset=start)
        ws, ed = np.isin(a, space), np.isin(a, edge)
        if (ws[1:] & ed[:-1]).any() or (ws[:-1] & ed[1:]).any() or (start == 0 and ws[0]):
            return True
    return bool(raw) and raw[-1] in space


def _clean_cell(v: Any) -> Any:
    if isinstance(v, str):
        v = v.strip()
        return v if v != "" else None
    if isinstance(v, dt.datetime) and v.time() == dt.time(0):
        return v
    return v


def iter_physical_rows(path: Path, sheet: str, start_row: int = 1, max_rows: int | None = None):
    """Yield (excel_row_number, [values]) starting at ``start_row`` (1-based)."""
    ext = path.suffix.lower()
    if ext in EXCEL_EXTS:
        fast = _calamine_rows(path, sheet, start_row, max_rows)
        if fast is not None:
            for i, row in enumerate(fast):
                yield start_row + i, [_clean_calamine_cell(v) for v in row]
            return
        wb = _open_workbook(path)
        try:
            if sheet not in wb.sheetnames:
                raise TabularError(f"Sheet '{sheet}' not found in workbook")
            ws = wb[sheet]
            max_row = start_row + max_rows - 1 if max_rows else None
            for i, row in enumerate(ws.iter_rows(min_row=start_row, max_row=max_row, values_only=True)):
                yield start_row + i, [_clean_cell(v) for v in row]
        finally:
            wb.close()
    elif ext in CSV_EXTS:
        text = _read_csv_text(path)
        reader = csv.reader(io.StringIO(text), delimiter=_csv_delimiter(path, text))
        for n, row in enumerate(reader, start=1):
            if n < start_row:
                continue
            if max_rows and n >= start_row + max_rows:
                break
            yield n, [_clean_cell(v) for v in row]
    else:
        raise TabularError(f"Unsupported tabular file type: {ext}")


def _jsonable(v: Any) -> Any:
    if v is None:
        return None
    if isinstance(v, float) and pd.isna(v):
        return None
    if isinstance(v, (dt.datetime, dt.date, dt.time)):
        return v.isoformat()
    if isinstance(v, (int, float, str, bool)):
        return v
    if isinstance(v, pd.Timestamp):
        return v.isoformat()
    return str(v)


def preview_rows(path: Path, sheet: str, max_rows: int = 30) -> dict:
    max_rows = min(max_rows, MAX_PREVIEW_ROWS)
    rows = [(n, vals) for n, vals in iter_physical_rows(path, sheet, 1, max_rows)]
    width = max((len(_trim(v)) for _, v in rows), default=0)
    return {
        "sheet": sheet,
        "width": width,
        "columns": [column_letter(i + 1) for i in range(width)],
        "rows": [{"row": n, "cells": [_jsonable(c) for c in (vals + [None] * width)[:width]]} for n, vals in rows],
        "suggested_header_row": suggest_header_row(rows),
    }


def _trim(vals: list) -> list:
    end = len(vals)
    while end and vals[end - 1] is None:
        end -= 1
    return vals[:end]


def column_letter(n: int) -> str:
    s = ""
    while n:
        n, r = divmod(n - 1, 26)
        s = chr(65 + r) + s
    return s


def suggest_header_row(rows: list[tuple[int, list]]) -> int | None:
    """Heuristic: first row that is mostly text, reasonably wide, and followed by data."""
    filled = [(n, [v for v in vals if v is not None], vals) for n, vals in rows]
    if not filled:
        return None
    widest = max(len(f) for _, f, _ in filled) or 1
    for idx, (n, nonempty, _) in enumerate(filled):
        if len(nonempty) < max(2, int(widest * 0.6)):
            continue
        text_share = sum(isinstance(v, str) for v in nonempty) / len(nonempty)
        if text_share < 0.8:
            continue
        if len(set(map(str, nonempty))) != len(nonempty):
            continue
        following = filled[idx + 1 : idx + 4]
        if following and any(len(f) >= len(nonempty) * 0.5 for _, f, _ in following):
            return n
    return filled[0][0]


@dataclass
class LoadedTable:
    df: pd.DataFrame
    warnings: list[str] = field(default_factory=list)
    header_row: int = 1
    first_data_row: int = 2
    original_headers: list[Any] = field(default_factory=list)
    letters: dict[str, str] = field(default_factory=dict)   # column name -> Excel column letter


def load_table(path: Path, sheet: str, header_row: int = 1, columns: list[str] | None = None,
               max_rows: int | None = None) -> LoadedTable:
    """Load a rectangular table whose header sits on 1-based ``header_row``."""
    if header_row < 1:
        raise TabularError("Header row must be 1 or greater (Excel row numbers start at 1)")
    warnings: list[str] = []
    csv_file = is_csv(path)
    limit = (max_rows + 1) if max_rows else None
    grid = _csv_grid(path, header_row, limit) if csv_file else None
    if grid is None:
        # Shorter rows are padded with None.
        grid = pd.DataFrame([vals for _, vals in iter_physical_rows(path, sheet, header_row, limit)], dtype=object)
    if not len(grid.index):
        raise TabularError(f"Row {header_row} is beyond the end of the sheet")

    header = [None if _isna(v) else v for v in grid.iloc[0].tolist()]
    body = grid.iloc[1:]
    filled = body.notna().to_numpy()
    # Drop trailing fully-empty rows.
    used = np.flatnonzero(filled.any(axis=1))
    end = used[-1] + 1 if len(used) else 0
    body, filled = body.iloc[:end], filled[:end]
    has_data = filled.any(axis=0)

    names: list[str] = []
    seen: dict[str, int] = {}
    keep: list[int] = []
    for i, h in enumerate(header):
        if h is None and not has_data[i]:
            continue
        name = str(h).strip() if h is not None else ""
        if isinstance(h, (dt.datetime, dt.date)):
            name = h.date().isoformat() if isinstance(h, dt.datetime) else h.isoformat()
        if not name:
            name = f"Column {column_letter(i + 1)}"
            warnings.append(f"Blank header in column {column_letter(i + 1)} renamed to '{name}'")
        if name in seen:
            seen[name] += 1
            new = f"{name} ({seen[name]})"
            warnings.append(f"Duplicate header '{name}' in column {column_letter(i + 1)} renamed to '{new}'")
            name = new
        else:
            seen[name] = 1
        names.append(name)
        keep.append(i)

    data = body.iloc[:, keep]
    blank = ~filled[:, keep].any(axis=1)
    if blank.any():
        warnings.append(f"{int(blank.sum())} blank row(s) inside the table were removed")
        data = data[~blank]
    df = data.set_axis(names, axis=1).reset_index(drop=True)
    df = type_csv_columns(df) if csv_file else normalise_types(df)

    letters = {n: column_letter(i + 1) for n, i in zip(names, keep)}
    t = LoadedTable(df=df, warnings=warnings, header_row=header_row, first_data_row=header_row + 1,
                    original_headers=[_jsonable(h) for h in header], letters=letters)
    return select_columns(t, columns)


def select_columns(t: LoadedTable, columns: list[str] | None) -> LoadedTable:
    """Narrow a loaded table to the selection's column subset (all columns when there is none)."""
    if columns:
        missing = [c for c in columns if c not in t.df.columns]
        if missing:
            raise TabularError(f"Selected column(s) no longer present: {', '.join(missing)}")
        t = replace(t, df=t.df[columns])
    return replace(t, warnings=list(t.warnings), letters={c: t.letters[c] for c in t.df.columns if c in t.letters})


def normalise_types(df: pd.DataFrame) -> pd.DataFrame:
    """Convert object columns to numeric/datetime only when every non-null value agrees."""
    out = {}
    for col in df.columns:
        s = df[col]
        nonnull = s.dropna()
        if nonnull.empty:
            out[col] = s
        elif all(isinstance(v, bool) for v in nonnull):
            out[col] = s.astype("boolean")
        elif all(isinstance(v, (int, float)) and not isinstance(v, bool) for v in nonnull):
            num = pd.to_numeric(s)
            if all(float(v).is_integer() for v in nonnull):
                num = num.astype("Int64")
            out[col] = num
        elif all(isinstance(v, (dt.datetime, dt.date)) for v in nonnull):
            out[col] = pd.to_datetime(s)
        else:
            out[col] = s.map(lambda v: None if v is None else (v if isinstance(v, str) else _jsonable(v))).astype(object)
    return pd.DataFrame(out, index=df.index)


# What a CSV cell must look like to be read as a number, date or true/false. Numbers are plain:
# no thousands separators, and no leading zeros, so codes like 00123 stay text; whole parts of at
# most 15 digits, so long IDs stay text rather than losing digits as JavaScript numbers in charts.
# Dates are ISO only: 03/04/2026 could be either month order.
_CSV_NUMBER = r"-?(?:(?:0|[1-9]\d{0,14})(?:\.\d+)?|\.\d+)(?:[eE][-+]?\d{1,3})?"
_CSV_DATE = r"\d{4}-\d{2}-\d{2}(?:[T ]\d{2}:\d{2}(?::\d{2}(?:\.\d{1,6})?)?)?"
_CSV_BOOL = {"true": True, "false": False}


def type_csv_columns(df: pd.DataFrame) -> pd.DataFrame:
    """Type a CSV's text columns as numbers, dates or true/false when every non-blank value is
    written as one, mirroring what a workbook's cells give ``normalise_types``. Each check runs
    on a column's distinct values, so repeated categories cost one test each."""
    out = {}
    for col in df.columns:
        s = df[col]
        blank = s.isna()
        values = pd.Series(pd.unique(s[~blank].to_numpy()), dtype=object)
        if values.empty:
            out[col] = s
        elif _all_match(_CSV_NUMBER, values):
            num = pd.to_numeric(s.astype(object))
            vals = num.dropna()
            if (vals.abs() < 2**53).all() and (vals == np.floor(vals)).all():
                num = num.astype("Int64")
            out[col] = num
        elif len(values) <= 8 and values.str.lower().isin(_CSV_BOOL).all():
            out[col] = s.str.lower().map(_CSV_BOOL).astype("boolean")
        elif _all_match(_CSV_DATE, values) and pd.to_datetime(values, format="ISO8601", errors="coerce").notna().all():
            out[col] = pd.to_datetime(s, format="ISO8601")
        else:
            out[col] = s.where(~blank, None) if blank.any() else s
    return pd.DataFrame(out, index=df.index)


def _all_match(pattern: str, values: pd.Series) -> bool:
    """Whether every value matches ``pattern`` in full: one regex run over the values joined by
    line breaks, many times quicker than a match per value."""
    joined = "\n".join(values)
    if joined.count("\n") != len(values) - 1:   # a value holds a line break, which no pattern allows
        return False
    return re.fullmatch(f"(?:{pattern})(?:\n(?:{pattern}))*", joined) is not None


def column_type(s: pd.Series) -> str:
    if s.dropna().empty:
        return "empty"
    if pd.api.types.is_bool_dtype(s):
        return "boolean"
    if pd.api.types.is_integer_dtype(s):
        return "integer"
    if pd.api.types.is_numeric_dtype(s):
        return "number"
    if pd.api.types.is_datetime64_any_dtype(s):
        return "date"
    return "text"


def profile(df: pd.DataFrame, sample_count: int = 3) -> list[dict]:
    cols = []
    for col in df.columns:
        s = df[col]
        nonnull = s.dropna()
        samples = []
        for v in nonnull.unique()[:sample_count]:
            jv = _jsonable(v.item() if hasattr(v, "item") else v)
            if isinstance(jv, str) and len(jv) > 60:
                jv = jv[:57] + "..."
            samples.append(jv)
        cols.append({
            "name": str(col),
            "type": column_type(s),
            "non_null": int(nonnull.shape[0]),
            "nulls": int(s.isna().sum()),
            "distinct": int(nonnull.nunique()),
            "samples": samples,
        })
    return cols


def schema_of(df: pd.DataFrame) -> list[dict]:
    return [{"name": str(c), "type": column_type(df[c])} for c in df.columns]


def is_csv(path: Path | str) -> bool:
    return Path(path).suffix.lower() in CSV_EXTS


def schema_drift(saved: list[dict] | None, current: list[dict], from_csv: bool = False) -> dict | None:
    """Compare a saved schema with the current one; None means no drift.

    ``from_csv``: CSV columns were all read as text before ``type_csv_columns``, so a schema saved
    then shows text where numbers or dates are now; that is not a change in the file."""
    if not saved:
        return None
    old = {c["name"]: c["type"] for c in saved}
    new = {c["name"]: c["type"] for c in current}
    missing = [c for c in old if c not in new]
    added = [c for c in new if c not in old]
    # "empty" columns carry no type information, so they never count as a type change.
    changed = [
        {"name": c, "was": old[c], "now": new[c]}
        for c in old
        if c in new and old[c] != new[c] and "empty" not in (old[c], new[c])
        and {old[c], new[c]} != {"integer", "number"}
        and not (from_csv and old[c] == "text")
    ]
    if not (missing or added or changed):
        return None
    return {"missing": missing, "added": added, "type_changes": changed}


def to_json_value(v: Any) -> Any:
    """Convert a pandas/numpy scalar into a JSON-friendly Python value."""
    if _isna(v):
        return None
    if isinstance(v, pd.Timestamp):
        return v.isoformat()
    return _jsonable(v.item() if hasattr(v, "item") and not isinstance(v, str) else v)


def column_values(s: pd.Series, limit: int | None = None) -> list[Any]:
    """One column as JSON-ready values, converted like ``to_json_value`` but a column at a time:
    fast enough for hundreds of thousands of rows. Infinite numbers become None."""
    if limit is not None:
        s = s.iloc[:limit]
    if pd.api.types.is_datetime64_any_dtype(s):
        if getattr(s.dt, "tz", None) is None and not (s.dt.microsecond.fillna(0) != 0).any():
            # numpy formats as YYYY-MM-DDTHH:MM:SS many times faster than strftime.
            out = np.datetime_as_string(s.to_numpy(dtype="datetime64[s]"), unit="s").astype(object)
            out[s.isna().to_numpy()] = None
            return out.tolist()
        return [None if pd.isna(v) else v.isoformat() for v in s]
    if pd.api.types.is_float_dtype(s):
        s = s.where(np.isfinite(s))
    if pd.api.types.is_numeric_dtype(s) or pd.api.types.is_bool_dtype(s):
        return s.astype(object).where(s.notna(), None).tolist()
    return [v if v is None or isinstance(v, str) else to_json_value(v) for v in s.tolist()]


def records(df: pd.DataFrame, limit: int | None = None) -> list[list[Any]]:
    sub = df if limit is None else df.head(limit)
    return [[to_json_value(v) for v in row] for row in sub.itertuples(index=False, name=None)]


def _isna(v: Any) -> bool:
    try:
        return bool(pd.isna(v))
    except (TypeError, ValueError):
        return False
