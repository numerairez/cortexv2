"""Writing tables to .xlsx/.csv without formula injection.

Text that starts with = + - @ (or tab/CR) would be interpreted as a formula by
Excel. In .xlsx we store such values explicitly as text cells; in .csv (where no
type information exists) we prefix them with an apostrophe.
"""

from __future__ import annotations

import csv
import datetime as dt
from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd

FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")


def _py(v: Any) -> Any:
    if v is None:
        return None
    try:
        if pd.isna(v):
            return None
    except (TypeError, ValueError):
        pass
    if isinstance(v, pd.Timestamp):
        return v.to_pydatetime()
    if isinstance(v, np.generic):
        return v.item()
    return v


def is_formula_like(v: Any) -> bool:
    return isinstance(v, str) and v.startswith(FORMULA_PREFIXES)


def write_csv(df: pd.DataFrame, path: Path) -> None:
    with path.open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow([("'" + str(c)) if is_formula_like(str(c)) else str(c) for c in df.columns])
        for row in df.itertuples(index=False, name=None):
            out = []
            for v in row:
                v = _py(v)
                if isinstance(v, (dt.datetime, dt.date)):
                    v = v.isoformat(sep=" ") if isinstance(v, dt.datetime) else v.isoformat()
                if is_formula_like(v):
                    v = "'" + v
                out.append("" if v is None else v)
            w.writerow(out)


def _write_sheet(ws, df: pd.DataFrame) -> None:
    from openpyxl.styles import Font

    ws.append([str(c) for c in df.columns])
    for cell in ws[1]:
        cell.font = Font(bold=True)
        if is_formula_like(cell.value):
            cell.data_type = "s"
    for r_idx, row in enumerate(df.itertuples(index=False, name=None), start=2):
        for c_idx, v in enumerate(row, start=1):
            v = _py(v)
            if isinstance(v, dt.datetime) and v.tzinfo is not None:
                v = v.replace(tzinfo=None)
            cell = ws.cell(row=r_idx, column=c_idx, value=v)
            if is_formula_like(v):
                cell.data_type = "s"  # store as literal text, never a formula
            elif isinstance(v, dt.datetime):
                cell.number_format = "yyyy-mm-dd" if v.time() == dt.time(0) else "yyyy-mm-dd hh:mm"
    ws.freeze_panes = "A2"
    for i, col in enumerate(df.columns, start=1):
        sample = [len(str(col))] + [len(str(_py(v) or "")) for v in df.iloc[:200, i - 1]]
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = min(60, max(8, max(sample) + 2))


def write_xlsx(df: pd.DataFrame, path: Path, summary_rows: list[tuple[str, Any]] | None = None,
               sheet_name: str = "Result") -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name[:31]
    _write_sheet(ws, df)
    if summary_rows:
        s = wb.create_sheet("Run summary")
        for key, value in summary_rows:
            s.append([key, value if not isinstance(value, (list, dict)) else str(value)])
            for cell in s[s.max_row]:
                if is_formula_like(cell.value):
                    cell.data_type = "s"
        for cell in s["A"]:
            cell.font = Font(bold=True)
        s.column_dimensions["A"].width = 28
        s.column_dimensions["B"].width = 100
    wb.save(path)
